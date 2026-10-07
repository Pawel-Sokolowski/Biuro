from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict


PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")


def generate_contract(
    template_path: str,
    output_dir: str,
    company_data: Dict[str, Any],
) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "Brak biblioteki python-docx. Zainstaluj zależności aplikacji."
        ) from exc

    template = Path(template_path)
    destination = Path(output_dir)
    if template.suffix.lower() != ".docx":
        raise ValueError("Szablon umowy musi być plikiem DOCX.")
    if not template.is_file():
        raise FileNotFoundError(f"Nie znaleziono szablonu: {template}")
    if not destination.is_dir():
        raise NotADirectoryError(f"Nie znaleziono folderu zapisu: {destination}")

    values = {key: "" if value is None else str(value) for key, value in company_data.items()}
    document = Document(str(template))

    for paragraph in _all_paragraphs(document):
        paragraph.text = _replace_placeholders(paragraph.text, values)

    company_name = _safe_filename(values.get("nazwa") or "firma")
    output_path = destination / f"Umowa - {company_name}.docx"
    document.save(str(output_path))
    return str(output_path)


def _replace_placeholders(text: str, values: Dict[str, str]) -> str:
    return PLACEHOLDER_PATTERN.sub(lambda match: values.get(match.group(1), match.group(0)), text)


def _all_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs

    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                yield paragraph
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def _safe_filename(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]', "_", value).strip().rstrip(".")
    return cleaned or "firma"
