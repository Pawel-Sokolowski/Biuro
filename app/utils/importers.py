import csv
import hashlib
import io
import os
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
import psycopg

from app.db.dao_audit import record_import
from app.db.dao_conn import get_conn
from app.utils.parsing import (
    to_float,
    normalize_identifier,
    normalize_pkd_codes,
    parse_date_value,
    resolve_employee_ids,
)


# ============================================================
# PARSERY POMOCNICZE
# ============================================================
def split_address_line(line: str) -> Tuple[str, str, str]:
    if not line:
        return "", "", ""

    text = str(line).strip()

    match = re.match(r"(.+?)\s+(\d+)\s*/\s*(\d+)", text)
    if match:
        return match.group(1).strip(), match.group(2).strip(), match.group(3).strip()

    match = re.match(r"(.+?)\s+(\d+)$", text)
    if match:
        return match.group(1).strip(), match.group(2).strip(), ""

    match = re.match(r"(\d+)\s*/\s*(\d+)", text)
    if match:
        return "", match.group(1).strip(), match.group(2).strip()

    match = re.match(r"(\d+)$", text)
    if match:
        return "", match.group(1).strip(), ""

    return text, "", ""


def parse_num(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    number = to_float(value, default=float("nan"))
    return None if number != number else number


PRICE_PATTERN = r"(\d{1,3}(?:\s\d{3})+(?:[.,]\d{1,2})?|\d+(?:[.,]\d{1,2})?)"


def _normalize_pricing_text(text: Any) -> str:
    normalized = str(text or "")
    normalized = normalized.replace("\xa0", " ")
    normalized = normalized.replace("\r", " ")
    normalized = normalized.replace("\n", " ")
    return re.sub(r"\s+", " ", normalized).strip()


def _price_to_float(value: str) -> float:
    return to_float(value)


def _find_price(pattern: str, text: str) -> Optional[float]:
    match = re.search(pattern, text, flags=re.IGNORECASE)
    if match is None:
        return None
    return _price_to_float(match.group(1))


_DOC_UNIT = r"(?:szt\.?|st\.?|dok\.?|dokument(?:ów|y)?)"


def parse_cennik_ksiegowy(text: str) -> Dict[str, Any]:
    out: Dict[str, Any] = {
        "limit_dokumentow": None,
        "cena_za_dokument": None,
        "ryczalt_ksiegowy": None,
        "zdrowotna": None,
    }

    if not text:
        return out

    normalized = _normalize_pricing_text(text)
    flags = re.IGNORECASE

    # „ryczałt” to stała opłata; „do 350 zł” oznacza cenę, nie limit dokumentów.
    limit_patterns = (
        rf"\b(?:ponad\s+limit|powyżej\s+limitu)\s+(\d+)\s*{_DOC_UNIT}",
        rf"(?<![\d.,])(\d+)\s*{_DOC_UNIT}(?!\w)",
        r"\blimit\s+(\d+)(?!\s*(?:zł|pln|[.,]\d))",
    )
    for pattern in limit_patterns:
        match = re.search(pattern, normalized, flags=flags)
        if match is not None:
            out["limit_dokumentow"] = int(match.group(1))
            break

    out["cena_za_dokument"] = _find_price(
        r"\b(?:ponad\s+limit|powyżej\s+limitu)\b\s*"
        rf"(?:\d+\s*{_DOC_UNIT}(?!\w)\s*)?"
        rf"\D*?{PRICE_PATTERN}\s*(?:zł|pln)?",
        normalized,
    )
    out["ryczalt_ksiegowy"] = _find_price(
        rf"\brycza[łl]t\b\s*(?:do\s+)?{PRICE_PATTERN}",
        normalized,
    )
    out["zdrowotna"] = _find_price(
        rf"\bzdrowotna\b\D*?{PRICE_PATTERN}",
        normalized,
    )
    if out["ryczalt_ksiegowy"] is None:
        # Zapis typu „260 zł + zdrowotna 60”: cena bazowa to pierwsza kwota poza limitem i zdrowotną.
        rest = re.sub(rf"\b(?:ponad\s+limit|powyżej\s+limitu)\b.*", " ", normalized, flags=flags)
        rest = re.sub(rf"\bzdrowotna\b\D*?{PRICE_PATTERN}", " ", rest, flags=flags)
        rest = re.sub(rf"(?<![\d.,])\d+\s*{_DOC_UNIT}(?!\w)", " ", rest, flags=flags)
        out["ryczalt_ksiegowy"] = _find_price(rf"(?<![\d.,]){PRICE_PATTERN}\s*(?:zł|pln)", rest)
    return out

def _find_service_price(text: str, service_pattern: str) -> Optional[float]:
    pattern = (
        rf"(?:usług(?:a|i)?\s*[-:]?\s*)?"
        rf"\b{service_pattern}\b"
        rf"\s*[-:]?\s*"
        rf"{PRICE_PATTERN}"
    )
    match = re.search(pattern, text, flags=re.IGNORECASE)
    if match is None:
        return None
    return _price_to_float(match.group(1))


def parse_cennik_kadrowy(text: str) -> Dict[str, Any]:
    out: Dict[str, Any] = {
        "zdrowotna": None,
        "uop_ilosc": 0,
        "uop_cennik": None,
        "uz_ilosc": 0,
        "uz_cennik": None,
        "uod_ilosc": 0,
        "uod_cennik": None,
    }

    if not text:
        return out

    normalized = _normalize_pricing_text(text)
    zdrowotna = _find_service_price(normalized, r"zdrowotna")
    uop = _find_service_price(normalized, r"u\s*o\s*p")
    uz = _find_service_price(normalized, r"u\s*z")
    uod = _find_service_price(normalized, r"u\s*o\s*d")

    if zdrowotna is not None:
        out["zdrowotna"] = zdrowotna
    if uop is not None:
        out["uop_ilosc"] = 1
        out["uop_cennik"] = uop
    if uz is not None:
        out["uz_ilosc"] = 1
        out["uz_cennik"] = uz
    if uod is not None:
        out["uod_ilosc"] = 1
        out["uod_cennik"] = uod

    return out


def parse_cennik_czasowy(text: str) -> Dict[str, Any]:
    out = {"ryczalt_czasowy": None}
    if not text:
        return out

    normalized = _normalize_pricing_text(text)
    match = re.search(rf"rycza[łl]t.*?{PRICE_PATTERN}", normalized, flags=re.IGNORECASE)
    if match:
        out["ryczalt_czasowy"] = _price_to_float(match.group(1))
    return out


# ============================================================
# WCZYTANIE PLIKU CSV/XLSX
# ============================================================
def _load_rows_from_file(filepath: str) -> Tuple[List[str], List[Any]]:
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".csv":
        with open(filepath, "rb") as handle:
            raw = handle.read()
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = raw.decode("cp1250")
        # Separator ustalamy po nagłówku; w danych przecinki są też w cenach (230,00).
        first_line = content.split("\n", 1)[0]
        delimiter = max((";", "\t", ","), key=first_line.count)
        reader = csv.reader(io.StringIO(content, newline=""), delimiter=delimiter)
        header = next(reader, None)
        rows = list(reader)
    else:
        workbook = openpyxl.load_workbook(filepath, data_only=True)
        worksheet = workbook.active
        header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        rows = list(worksheet.iter_rows(min_row=2, values_only=True))

    if not header:
        raise ValueError("Brak nagłówków w pliku.")

    headers = [str(value).strip() if value is not None else "" for value in header]
    return headers, rows


def import_tax_offices_from_excel(filepath: str) -> int:
    """Importuje lub aktualizuje katalog urzędów skarbowych z XLSX/CSV."""
    headers, rows = _load_rows_from_file(filepath)

    def normalize_header(value: str) -> str:
        return re.sub(r"\s+", " ", value).strip().upper()

    index = {normalize_header(header): pos for pos, header in enumerate(headers)}

    def value(row: Any, *names: str) -> str:
        for name in names:
            pos = index.get(normalize_header(name))
            if pos is not None and pos < len(row):
                return str(row[pos] or "").strip()
        return ""

    records = []
    for row in rows:
        name = value(row, "NAZWA URZĘDU")
        if name:
            records.append((
                value(row, "WOJEWÓDZTWO"), value(row, "TYP"), name,
                value(row, "KOD"), value(row, "MIASTO"), value(row, "ULICA"),
                value(row, "NR BUDYNKU / LOKALU"), value(row, "POCZTA"),
                value(row, "NR TELEFONU wraz z nr kierunkowym", "NR TELEFONU"),
                value(row, "NR FAXU"), value(row, "ADRES E-MAIL URZĘDU"),
                value(row, "KOD JEDNOSTKI"), value(row, "ADRESY STRON BIP JEDNOSTEK KAS"),
            ))

    if not records:
        raise ValueError("Nie znaleziono wierszy urzędów skarbowych.")

    conn = get_conn()
    try:
        conn.cursor().executemany(
            """
            INSERT INTO urzedy_skarbowe (
                wojewodztwo, typ, nazwa, kod, miasto, ulica, nr_budynku_lokalu,
                poczta, nr_telefonu, nr_faxu, email, kod_jednostki, adresy_bip
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT(kod_jednostki) DO UPDATE SET
                wojewodztwo = excluded.wojewodztwo,
                typ = excluded.typ,
                nazwa = excluded.nazwa,
                kod = excluded.kod,
                miasto = excluded.miasto,
                ulica = excluded.ulica,
                nr_budynku_lokalu = excluded.nr_budynku_lokalu,
                poczta = excluded.poczta,
                nr_telefonu = excluded.nr_telefonu,
                nr_faxu = excluded.nr_faxu,
                email = excluded.email,
                adresy_bip = excluded.adresy_bip
            """,
            records,
        )
        conn.commit()
        return len(records)
    finally:
        conn.close()


_TAX_OFFICE_STOPWORDS = {"urzad", "skarbowy", "us", "w", "we", "im", "nr", "i"}


def _office_tokens(text: str) -> List[str]:
    folded = re.sub(r"[^a-z0-9]+", " ", _fold_pl(text))
    return [t for t in folded.split() if t not in _TAX_OFFICE_STOPWORDS]


def _tokens_similar(a: str, b: str) -> bool:
    if a == b:
        return True
    n = min(len(a), len(b))
    return n >= 4 and a[: n - 1] == b[: n - 1]


def match_tax_office(value: Any) -> Optional[str]:
    """Dopasowuje nazwę z pliku do katalogu urzędów (wielkość liter, polskie znaki,
    odmiana miasta, skróty „US”); gdy brak dopasowania, zostawia wpisaną nazwę."""
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if not text:
        return None
    try:
        from app.db.dao_companies import get_tax_offices
        names = [
            re.sub(r"\s+", " ", str(o.get("nazwa") or "")).strip()
            for o in get_tax_offices()
        ]
        names = [n for n in names if n]
    except Exception:
        return text

    for name in names:
        if name.casefold() == text.casefold():
            return name

    wanted = _office_tokens(text)
    if not wanted:
        return text
    best, best_score, tie = None, 0.0, False
    for name in names:
        tokens = _office_tokens(name)
        if not tokens:
            continue
        hits = sum(1 for w in wanted if any(_tokens_similar(w, t) for t in tokens))
        if hits < len(wanted):
            continue
        score = hits / len(tokens)
        if score > best_score:
            best, best_score, tie = name, score, False
        elif score == best_score:
            tie = True
    if best and not tie:
        return best
    return text


# ============================================================
# BUDOWANIE PAYLOADU Z WIERSZA
# ============================================================
def _build_payload_from_row(
    headers: List[str],
    header_index: Dict[str, int],
    row: Any,
) -> Dict[str, Any]:
    del headers

    def _norm(name: str) -> str:
        return re.sub(r"\s+", " ", str(name)).strip().casefold()

    normalized_index = {_norm(k): v for k, v in header_index.items()}

    def get(col: str) -> Any:
        _CONSUMED_HEADERS.add(_norm(col))
        idx = normalized_index.get(_norm(col))
        value = row[idx] if idx is not None and idx < len(row) else None
        if isinstance(value, str):
            value = value.strip()
        elif isinstance(value, float) and value.is_integer():
            value = int(value)
        if isinstance(value, int) and not isinstance(value, bool):
            value = str(value)
        return value

    payload: Dict[str, Any] = {}

    payload["nazwa"] = get("Nazwa Firmy")
    payload["nazwa_skrocona"] = get("Nazwa skrócona")
    payload["nip"] = normalize_identifier(get("NIP"), {10})
    payload["regon"] = normalize_identifier(get("REGON"), {9, 14})
    payload["krs"] = normalize_identifier(get("KRS"), {10})
    payload["pesel"] = normalize_identifier(get("PESEL"), {11})

    payload["forma_dzialalnosci"] = get("Forma działalności")
    payload["forma_opodatkowania"] = get("Forma opodatkowania")
    payload["konto_bankowe"] = get("Rachunek bankowy") or get("Numer konta")
    payload["urzad_skarbowy"] = match_tax_office(get("Urząd Skarbowy") or get("Urząd skarbowy"))

    payload["analityka_fk"] = get("Analityka FK")
    payload["osoba_kontaktowa"] = get("Osoba kontaktowa")
    payload["adres_email"] = get("Adres email")
    payload["nr_telefonu"] = get("Nr telefonu")

    vat_raw = get("VAT czynny")
    if vat_raw in (None, ""):
        vat_raw = get("VAT")
    if vat_raw in (None, ""):
        vat_raw = get("Status VAT")
    if vat_raw not in (None, ""):
        active = _is_vat_active(vat_raw)
        payload["vat_czynny"] = 1 if active else 0
        payload["status_vat"] = "Czynny" if active else "Zwolniony"

    payload["miasto"] = get("Miasto")
    payload["kod_pocztowy"] = get("Kod pocztowy")

    (
        payload["ulica"],
        payload["adres_budynek"],
        payload["adres_lokal"],
    ) = split_address_line(get("Ulica"))

    payload["adres_zam_miasto"] = get("Adres zamieszkania - miasto")
    payload["adres_zam_kod_pocztowy"] = get("Adres zamieszkania - kod pocztowy")

    (
        payload["adres_zam_ulica"],
        payload["adres_zam_budynek"],
        payload["adres_zam_lokal"],
    ) = split_address_line(
        get("Adres zamieszkania - ulica")
    )

    payload["adres_koresp_miasto"] = (
        get("Adres korespondencyjny - miasto")
        or get("Adres korespondpondencyjny - miasto")
    )

    payload["adres_koresp_kod_pocztowy"] = get(
        "Adres korespondencyjny - kod pocztowy"
    )

    (
        payload["adres_koresp_ulica"],
        payload["adres_koresp_budynek"],
        payload["adres_koresp_lokal"],
    ) = split_address_line(
        get("Adres korespondencyjny - ulica")
    )

    payload["data_rozpoczecia_dg"] = parse_date_value(get("Data rozpoczęcia działalności gospodarczej"))
    payload["data_rozpoczecia_sprzedazy"] = parse_date_value(get("Data rozpoczęcia sprzedaży dla osoby fizycznej"))
    payload["data_koniec_zus_start"] = parse_date_value(get("Data zakończenia ulgi ZUS na start"))
    payload["data_koniec_zus_pref"] = parse_date_value(get("Data zakończenia ZUS preferencyjny"))
    payload["data_koniec_maly_zus"] = parse_date_value(get("Data zakończenia małego ZUS PLUS"))
    payload["data_koniec_a1"] = parse_date_value(get("Data zakończenia A1"))
    payload["data_rozpoczecia_obslugi"] = parse_date_value(get("Data rozpoczęcia obsługi"))
    payload["data_zakonczenia_obslugi"] = parse_date_value(get("Data zakończenia obsługi"))
    payload["zawieszenia_dzialalnosci"] = get("Zawieszenia działalności")
    payload["kody_pkd"] = normalize_pkd_codes(get("Kody PKD"))

    # --------------------------------------------------
    # KSIĘGOWOŚĆ
    # --------------------------------------------------

    cennik_ksiegowy_text = str(
        get("Cenniki usług księgowych") or ""
    )

    cks = parse_cennik_ksiegowy(
        cennik_ksiegowy_text
    )

    payload["cennik_ksiegowy"] = cennik_ksiegowy_text
    payload["cennik_ksiegowy_parsed"] = cks
    payload["limit_dokumentow"] = cks["limit_dokumentow"]
    payload["cena_netto_ksiegowosc"] = cks["ryczalt_ksiegowy"]
    payload["cena_ponad_limit"] = cks["cena_za_dokument"]
    payload["glowny_cennik"] = suggest_accounting_cennik_key(payload)

    # --------------------------------------------------
    # KADRY
    # --------------------------------------------------

    cennik_kadrowy_text = str(
        get("Cenniki usług kadrowych") or ""
    )
    ckad = parse_cennik_kadrowy(
        cennik_kadrowy_text
    )

    payload["cennik_kadrowy"] = cennik_kadrowy_text
    payload["cennik_kadrowy_parsed"] = ckad

    # --------------------------------------------------
    # ZDROWOTNA
    # --------------------------------------------------

    ksiegowosc_zdrowotna = cks.get("zdrowotna")
    kadry_zdrowotna = ckad.get("zdrowotna")
    payload["zdrowotna"] = (
        ksiegowosc_zdrowotna
        if ksiegowosc_zdrowotna is not None
        else kadry_zdrowotna
    )

    payload["wliczac_zdrowotna"] = (1 if ksiegowosc_zdrowotna is not None else 0)

    payload["uop_ilosc"] = ckad["uop_ilosc"]
    payload["uop_cennik"] = ckad["uop_cennik"]
    payload["uz_ilosc"] = ckad["uz_ilosc"]
    payload["uz_cennik"] = ckad["uz_cennik"]
    payload["uod_ilosc"] = ckad["uod_ilosc"]
    payload["uod_cennik"] = ckad["uod_cennik"]

    # --------------------------------------------------
    # CZASOWE
    # --------------------------------------------------

    cennik_czasowy_text = str(
        get("Cenniki czasowe") or ""
    )
    cczas = parse_cennik_czasowy(cennik_czasowy_text)
    payload["cennik_czasowy"] = cennik_czasowy_text
    payload["cennik_czasowy_parsed"] = cczas
    payload["ryczalt_czasowy"] = cczas["ryczalt_czasowy"]
    payload["assigned_raw"] = get("Przypisani pracownicy")

    return payload
# ============================================================
# IMPORT DO PAYLOADÓW (DLA WIZARDU)
# ============================================================
def import_companies_from_excel_file_payloads(filepath: str) -> Tuple[List[Dict[str, Any]], List[str]]:
    try:
        headers, rows = _load_rows_from_file(filepath)
    except Exception as exc:
        return [], [f"Plik nie mógł zostać wczytany: {exc}"]

    header_index = {header: idx for idx, header in enumerate(headers)}
    payloads: List[Dict[str, Any]] = []
    warnings: List[str] = []

    for row_number, row in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in row):
            continue

        payload = _build_payload_from_row(headers, header_index, row)
        assigned_raw = payload.get("assigned_raw")
        if assigned_raw:
            unresolved = [
                part.strip()
                for part in str(assigned_raw).split(",")
                if part.strip() and not re.match(r"^[A-Za-zĄąĆćĘęŁłŃńÓóŚśŹźŻż ]+$", part.strip())
            ]
            if unresolved:
                warnings.append(f"Wiersz {row_number}: nie znaleziono pracowników {', '.join(unresolved)}.")

        payloads.append(payload)

    known = set(_CONSUMED_HEADERS)
    unknown = [h for h in headers if h and re.sub(r"\s+", " ", h).strip().casefold() not in known]
    if payloads and unknown:
        warnings.append("Nie rozpoznano kolumn (dane z nich nie zostaną zaimportowane): " + ", ".join(unknown))

    return payloads, warnings


# ============================================================
# ZAPIS DO DB
# ============================================================
def _pricing_key(prefix: str, values: Dict[str, Any]) -> str:
    raw = "|".join(f"{key}={values.get(key)!r}" for key in sorted(values))
    digest = hashlib.sha1(raw.encode("utf-8")).hexdigest()[:12].upper()
    return f"IMPORT_{prefix}_{digest}"


_CONSUMED_HEADERS: set = set()

_VAT_ACTIVE_VALUES = {
    "tak", "t", "true", "yes", "y", "x", "1", "czynny", "vat czynny", "czynna",
    "miesieczny", "miesiecznie", "kwartalny", "kwartalnie", "m", "k",
}


def _is_vat_active(value: Any) -> bool:
    if value in (1, True):
        return True
    return _fold_pl(value).rstrip(".:;") in _VAT_ACTIVE_VALUES


def _fold_pl(value: Any) -> str:
    text = str(value or "").strip().lower().replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text)
    return text.translate(str.maketrans({
        "ą": "a", "ć": "c", "ę": "e", "ł": "l", "ń": "n",
        "ó": "o", "ś": "s", "ź": "z", "ż": "z",
    }))


# Prefiks klucza cennika: z formy opodatkowania, nie z tekstu "Ryczałt 1300"
ACCOUNTING_CENNIK_PREFIXES = {
    "KPIR": ("KPIR",),
    "RYCZALT": ("RYCZALT", "Ryczałt", "Ryczalt"),
    "KH": ("KH", "KR"),
}


def resolve_accounting_cennik_typ(
    forma_opodatkowania: Any = None,
    forma_dzialalnosci: Any = None,
) -> str:
    tax = _fold_pl(forma_opodatkowania)
    activity = _fold_pl(forma_dzialalnosci)

    if tax:
        if "ryczalt" in tax:
            return "RYCZALT"
        if "liniow" in tax or "skala" in tax or "kpir" in tax:
            return "KPIR"
        if "cit" in tax or "ksieg" in tax or tax in {"kh", "kr"}:
            return "KH"

    if "kpir" in activity or "osoba fizyczna" in activity:
        return "KPIR"
    if "cywiln" in activity:
        return "KPIR"
    if activity:
        return "KH"
    return "KPIR"


def _accounting_cennik_key_candidates(
    typ: str,
    vat_flag: str,
    limit: int,
    price: Any = None,
) -> List[str]:
    prefixes = ACCOUNTING_CENNIK_PREFIXES.get(typ, (typ,))

    if price is not None:
        price_key = _format_price_key(price)
        return [
            f"{prefix}_{vat_flag}_{limit}_{price_key}"
            for prefix in prefixes
        ]

    return [
        f"{prefix}_{vat_flag}_{limit}"
        for prefix in prefixes
    ]


def _accounting_cennik_spec(payload: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    parsed = payload.get("cennik_ksiegowy_parsed") or {}
    text = str(payload.get("cennik_ksiegowy") or "").strip()
    if not text:
        return None

    limit_price = parsed.get("cena_za_dokument")
    base_price = parsed.get("ryczalt_ksiegowy")
    if limit_price is None and base_price is None:
        return None

    typ = resolve_accounting_cennik_typ(
        payload.get("forma_opodatkowania"),
        payload.get("forma_dzialalnosci"),
    )

    status_vat = _fold_pl(payload.get("status_vat"))
    if status_vat in {"czynny", "zwolniony"}:
        vat_active = status_vat == "czynny"
    else:
        vat_active = _is_vat_active(payload.get("vat_czynny"))
    vat_flag = "T" if vat_active else "N"
    limit = parsed.get("limit_dokumentow") or 0

    # Cena podstawowa jest częścią klucza, np. KPIR_T_10_300.
    candidates = _accounting_cennik_key_candidates(typ, vat_flag, limit, base_price)

    return {
        "typ": typ,
        "candidates": candidates,
        "base_price": base_price,
        "limit_price": limit_price,
        "zdrowotna": parsed.get("zdrowotna") or payload.get("zdrowotna") or 0,
    }


def suggest_accounting_cennik_key(payload: Dict[str, Any]) -> Optional[str]:
    spec = _accounting_cennik_spec(payload)
    return spec["candidates"][0] if spec else None


def _ensure_accounting_cennik(cursor: psycopg.Cursor, payload: Dict[str, Any]) -> Optional[str]:
    spec = _accounting_cennik_spec(payload)
    if spec is None:
        return None

    candidates = spec["candidates"]
    base_price = spec["base_price"]
    # „netto” w cenniku to cena księgowości wraz ze zdrowotną.
    netto = round((base_price or 0) + float(spec["zdrowotna"] or 0), 2)
    placeholders = ", ".join("%s" for _ in candidates)

    cursor.execute(
        f"SELECT klucz FROM cenniki WHERE klucz IN ({placeholders}) LIMIT 1",
        candidates,
    )
    row = cursor.fetchone()
    if row:
        return row[0]

    generated_key = candidates[0]

    cursor.execute(
        """
        INSERT INTO cenniki (
            klucz, typ, cena_netto_ksiegowosc, cena_brutto_ksiegowosc,
            cena_za_dokument_ponad_limit, zdrowotna, netto, brutto
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        """,
        (
            generated_key,
            spec["typ"],
            base_price or 0,
            round((base_price or 0) * 1.23, 2),
            spec["limit_price"] or 0,
            spec["zdrowotna"],
            netto,
            round(netto * 1.23, 2),
        ),
    )

    return generated_key


def _decimal_price(value: Any) -> Decimal:
    if isinstance(value, Decimal):
        return value

    normalized = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".").strip()
    try:
        return Decimal(normalized)
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise ValueError(f"Nieprawidłowa cena cennika kadrowego: {value!r}") from exc


def _format_price_key(value: Any) -> str:
    price = _decimal_price(value).normalize()
    if price == price.to_integral():
        return str(int(price))
    return format(price, "f").rstrip("0").rstrip(".").replace(".", "_")


def _ensure_employee_cenniki(cursor: psycopg.Cursor, payload: Dict[str, Any]) -> Dict[str, str]:
    text = str(payload.get("cennik_kadrowy") or "").strip()
    if not text:
        return {}

    parsed = payload.get("cennik_kadrowy_parsed")
    if not isinstance(parsed, dict):
        parsed = {}

    parsed_from_text = parse_cennik_kadrowy(text)
    for field, value in parsed_from_text.items():
        if value is not None:
            parsed[field] = value

    service_definitions = (
        ("ZDROWOTNA", "Zdrowotna", parsed.get("zdrowotna")),
        ("UOP", "UoP", parsed.get("uop_cennik")),
        ("UZ", "UZ", parsed.get("uz_cennik")),
        ("UOD", "UoD", parsed.get("uod_cennik")),
    )

    keys: Dict[str, str] = {}
    for service_code, key_prefix, raw_price in service_definitions:
        if raw_price is None:
            continue

        price_netto = _decimal_price(raw_price).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        price_brutto = (price_netto * Decimal("1.23")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        key = f"{key_prefix}_{_format_price_key(price_netto)}"

        cursor.execute(
            """
            INSERT INTO cenniki_kadry (
                klucz, cena_netto_kadry, cena_brutto_kadry, opis
            ) VALUES (%s, %s, %s, %s)
            ON CONFLICT(klucz) DO UPDATE SET
                cena_netto_kadry = excluded.cena_netto_kadry,
                cena_brutto_kadry = excluded.cena_brutto_kadry,
                opis = excluded.opis,
                updated_at = to_char(now(), 'YYYY-MM-DD HH24:MI:SS')
            """,
            (key, float(price_netto), float(price_brutto), f"Import z Excela: {service_code}"),
        )
        keys[service_code] = key

    return keys


def _assign_employees(
    cursor: psycopg.Cursor,
    firma_id: int,
    assigned_raw: Any,
    warnings: List[str],
    row_number: int,
) -> None:
    if not assigned_raw:
        return

    assigned_ids, unresolved = resolve_employee_ids(cursor, assigned_raw)
    for employee_id in assigned_ids:
        cursor.execute(
            "INSERT INTO pracownik_firma (pracownik_id, firma_id, rola_id) VALUES (%s, %s, NULL) ON CONFLICT DO NOTHING",
            (employee_id, firma_id),
        )
    if unresolved:
        warnings.append(f"Wiersz {row_number}: nie znaleziono pracowników {', '.join(unresolved)}.")


# ============================================================
# ZAPIS DO DB
# ============================================================
def _apply_payload_to_db(
    cursor: psycopg.Cursor,
    payload: Dict[str, Any],
    row_number: int,
    warnings: List[str],
) -> bool:
    company_name = str(payload.get("nazwa", "") or "").strip()
    if not company_name:
        warnings.append(f"Wiersz {row_number}: pominięto pustą nazwę firmy.")
        return False

    from app.db.dao_companies import FIRMA_FIELDS, _apply_aliases

    # Kreator przekazuje klucze formularza (dz_*, ko_*, zam_*, kont_*) – mapujemy je na kolumny.
    payload.update({k: v for k, v in _apply_aliases(payload).items() if k not in payload})
    if payload.get("status_vat") in ("Czynny", "Zwolniony") and payload.get("vat_czynny") in (None, ""):
        payload["vat_czynny"] = 1 if payload["status_vat"] == "Czynny" else 0

    normalized_nip = normalize_identifier(payload.get("nip"), {10})
    company_id = None

    if normalized_nip:
        cursor.execute("SELECT id FROM firmy WHERE nip = %s", (normalized_nip,))
        row_db = cursor.fetchone()
        company_id = row_db[0] if row_db else None

    if company_id is None:
        cursor.execute("SELECT id FROM firmy WHERE nazwa = %s", (company_name,))
        row_db = cursor.fetchone()
        company_id = row_db[0] if row_db else None

    accounting_key = _ensure_accounting_cennik(cursor, payload)
    employee_keys = _ensure_employee_cenniki(cursor, payload)

    if "UOP" in employee_keys:
        payload["uop_cennik"] = employee_keys["UOP"]
    if "UZ" in employee_keys:
        payload["uz_cennik"] = employee_keys["UZ"]
    if "UOD" in employee_keys:
        payload["uod_cennik"] = employee_keys["UOD"]

    company_fields = [
        "nazwa", "nazwa_skrocona", "nip", "regon", "krs", "pesel",
        "forma_dzialalnosci", "forma_opodatkowania", "urzad_skarbowy", "konto_bankowe", "analityka_fk",
        "osoba_kontaktowa", "adres_email", "nr_telefonu", "miasto",
        "kod_pocztowy", "ulica", "adres_budynek", "adres_lokal",
        "adres_zam_miasto", "adres_zam_kod_pocztowy", "adres_zam_ulica",
        "adres_zam_budynek", "adres_zam_lokal", "adres_koresp_miasto",
        "adres_koresp_kod_pocztowy", "adres_koresp_ulica",
        "adres_koresp_budynek", "adres_koresp_lokal",
        "data_rozpoczecia_dg", "data_rozpoczecia_sprzedazy",
        "data_koniec_zus_start", "data_koniec_zus_pref",
        "data_koniec_maly_zus", "data_koniec_a1",
        "data_rozpoczecia_obslugi", "data_zakonczenia_obslugi",
        "zawieszenia_dzialalnosci", "kody_pkd", "cennik_ksiegowy",
        "cennik_kadrowy", "cennik_czasowy", "limit_dokumentow",
        "zdrowotna", "uop_ilosc", "uop_cennik", "uz_ilosc",
        "uz_cennik", "uod_ilosc", "uod_cennik", "wliczac_zdrowotna",
    ]

    cursor.execute(
        "SELECT column_name FROM information_schema.columns "
        "WHERE table_schema = current_schema() AND table_name = 'firmy'"
    )
    existing_columns = {row[0] for row in cursor.fetchall()}
    for field in FIRMA_FIELDS:
        if field in existing_columns and field not in company_fields and field not in (
            "glowny_cennik", "glowny_cennik_manualny",
        ):
            company_fields.append(field)

    if accounting_key:
        company_fields.append("glowny_cennik")
        payload["glowny_cennik"] = accounting_key

    fields = []
    values = []

    for field in company_fields:
        value = payload.get(field)
        if value not in (None, "") or field == "nazwa":
            fields.append(field)
            values.append(value)

    if company_id:
        assignments = ", ".join(f"{field} = %s" for field in fields)
        cursor.execute(
            f"UPDATE firmy SET {assignments} WHERE id = %s",
            (*values, company_id),
        )
        created = False
    else:
        placeholders = ", ".join("%s" for _ in fields)
        cursor.execute(
            f"INSERT INTO firmy ({', '.join(fields)}) VALUES ({placeholders}) RETURNING id",
            values,
        )
        _new_id = cursor.fetchone()[0]
        company_id = int(_new_id)
        created = True

    _ensure_contact(cursor, company_id, payload)
    _assign_employees(cursor, company_id, payload.get("assigned_raw"), warnings, row_number)
    return created


def _ensure_contact(cursor: psycopg.Cursor, firma_id: int, payload: Dict[str, Any]) -> None:
    """Osoba kontaktowa z pliku trafia też do kontaktów firmy (z nich powstają filtry e-mail)."""
    email = str(payload.get("adres_email") or "").strip()
    person = str(payload.get("osoba_kontaktowa") or "").strip()
    phone = str(payload.get("nr_telefonu") or "").strip()
    if not (email or person):
        return
    if email:
        exists = cursor.execute(
            "SELECT 1 FROM company_contacts WHERE firma_id = %s AND lower(email) = %s",
            (firma_id, email.lower()),
        ).fetchone()
    else:
        exists = cursor.execute(
            "SELECT 1 FROM company_contacts WHERE firma_id = %s AND (imie || ' ' || COALESCE(nazwisko, '')) ILIKE %s",
            (firma_id, f"{person}%"),
        ).fetchone()
    if exists:
        return
    first, _, last = person.partition(" ")
    cursor.execute(
        "INSERT INTO company_contacts (firma_id, imie, nazwisko, email, telefon, rola) VALUES (%s, %s, %s, %s, %s, %s)",
        (firma_id, first, last.strip(), email, phone, ""),
    )


# ============================================================
# IMPORTERY FIRM
# ============================================================
def import_companies_from_excel_file(filepath: str, db_path: Optional[str] = None) -> Dict[str, Any]:
    try:
        headers, rows = _load_rows_from_file(filepath)
    except Exception as exc:
        return {"created": 0, "updated": 0, "warnings": [f"Plik nie mógł zostać wczytany: {exc}"]}

    header_index = {header: idx for idx, header in enumerate(headers)}
    conn = get_conn()
    cursor = conn.cursor()
    created_count = 0
    updated_count = 0
    warnings: List[str] = []

    try:
        for row_number, row in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in row):
                continue
            payload = _build_payload_from_row(headers, header_index, row)
            if _apply_payload_to_db(cursor, payload, row_number, warnings):
                created_count += 1
            else:
                updated_count += 1
            if row_number % 20 == 0:
                conn.commit()

        conn.commit()
        result = {"created": created_count, "updated": updated_count, "warnings": warnings}
        record_import("firmy", filepath, result)
        return result
    finally:
        conn.close()


def save_imported_payload(
    payload: Dict[str, Any],
    db_path: Optional[str] = None,
    row_number: int = 2,
) -> Dict[str, Any]:
    """Zapisuje jeden zatwierdzony payload z kreatora importu."""
    conn = get_conn()
    cursor = conn.cursor()
    warnings: List[str] = []
    try:
        created = _apply_payload_to_db(cursor, payload, row_number, warnings)
        conn.commit()
        return {"created": int(created), "updated": int(not created), "warnings": warnings}
    finally:
        conn.close()


def import_companies_fast_from_excel_file(filepath: str, db_path: Optional[str] = None) -> Dict[str, Any]:
    try:
        headers, rows = _load_rows_from_file(filepath)
    except Exception as exc:
        return {"created": 0, "updated": 0, "warnings": [f"Plik nie mógł zostać wczytany: {exc}"]}

    header_index = {header: idx for idx, header in enumerate(headers)}
    conn = get_conn()
    cursor = conn.cursor()
    created = 0
    updated = 0
    warnings: List[str] = []

    try:
        for row_number, row in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in row):
                continue
            payload = _build_payload_from_row(headers, header_index, row)
            if _apply_payload_to_db(cursor, payload, row_number, warnings):
                created += 1
            else:
                updated += 1

        conn.commit()
        result = {"created": created, "updated": updated, "warnings": warnings}
        record_import("firmy_fast", filepath, result)
        return result
    finally:
        conn.close()


# ============================================================
# IMPORT CENNIKA KSIĘGOWEGO
# ============================================================
def build_pricing_key(base_key: str, price: Any) -> str:
    """Dokleja cenę podstawową do klucza: KPIR_T_10 + 300 -> KPIR_T_10_300."""
    key = str(base_key or "").strip()
    if not key or not price:
        return key
    price_key = _format_price_key(price)
    parts = re.split(r"[_\-]", key)
    if len(parts) >= 4 and parts[-1] == price_key:
        return key
    return f"{key}_{price_key}"


def import_cennik_from_excel(filepath: str, db_path: Optional[str] = None) -> List[Dict[str, Any]]:
    del db_path
    try:
        headers, rows = _load_rows_from_file(filepath)
    except Exception:
        return []

    normalized_headers = [header.lower() for header in headers]
    idx = {header: pos for pos, header in enumerate(normalized_headers)}

    required = [
        "klucz",
        "cena netto 2024",
        "wzrost ilości dokumentów",
        "cena za dokument",
        "zdrowotna",
        "netto",
        "brutto",
    ]

    if any(column not in idx for column in required):
        return []

    def get_cell(row: Any, key: str) -> Any:
        position = idx.get(key)
        return row[position] if position is not None and position < len(row) else None

    payloads = []

    for row in rows:
        klucz_raw = get_cell(row, "klucz")

        if klucz_raw is None or not str(klucz_raw).strip():
            continue
        # Klucz jest automatycznie rozszerzany o wartość z kolumny Netto.
        base_key = str(klucz_raw).strip()

        netto = parse_num(get_cell(row, "netto")) or 0.0

        generated_key = f"{base_key}_{_format_price_key(netto)}"

        payloads.append({
            "klucz": generated_key,
            "cena_netto_ksiegowosc": parse_num(get_cell(row, "cena netto 2024")) or 0.0,
            "cena_za_dokument_ponad_limit": parse_num(get_cell(row, "cena za dokument")) or 0.0,
            "zdrowotna": parse_num(get_cell(row, "zdrowotna")) or 0.0,
            "netto": netto,
            "brutto": parse_num(get_cell(row, "brutto")) or 0.0,
            "wzrost_ilosci_dokumentow": parse_num(get_cell(row, "wzrost ilości dokumentów")) or 0.0,
        })

    return payloads

def import_ceny_za_dokument_from_excel(filepath: str, db_path: Optional[str] = None) -> List[Dict[str, Any]]:
    del db_path
    try:
        headers, rows = _load_rows_from_file(filepath)
    except Exception:
        return []

    normalized_headers = [header.lower() for header in headers]
    idx = {header: pos for pos, header in enumerate(normalized_headers)}
    if "klucz" not in idx or "cena" not in idx:
        return []

    payloads = []
    for row in rows:
        if not any(cell not in (None, "") for cell in row):
            continue
        klucz = str(row[idx["klucz"]] or "").strip()
        if not klucz:
            continue

        cena = parse_num(row[idx["cena"]])
        typ = klucz
        vat = 0
        if "_" in klucz:
            parts = klucz.split("_")
            typ = parts[0]
            vat = int(parts[-1]) if parts[-1].isdigit() else 0

        payloads.append({"typ": typ, "vat": vat, "cena": cena})
    return payloads