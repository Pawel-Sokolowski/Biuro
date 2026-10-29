"""Korekty nigdy nie trafiają do szablonu zwykłych faktur Saldeo."""
import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from app.utils import saldeo_exporter as exporter
from tests.test_saldeo_export_names import FakeConnection


class SaldeoCorrectionTests(unittest.TestCase):
    def setUp(self):
        self.invoice = {
            "id": 10, "firma_id": 1, "rodzaj": "faktura", "numer": "FV/001/09/2026",
            "miesiac": "2026-09", "kwota_netto": 100, "kwota_brutto": 123,
            "pozycje_json": [{"nazwa": "Usługa", "ilosc": 1, "cena_netto": 100,
                              "netto": 100, "brutto": 123, "vat": "23%"}],
        }
        self.correction = dict(
            self.invoice, id=11, rodzaj="korekta", numer="KOR/001/09/2026",
            korygowana_faktura_id=10, korygowana_faktura_numer="FV/001/09/2026",
            kwota_netto=-100, kwota_brutto=-123,
            pozycje_json=json.dumps([{"nazwa": "Usługa", "ilosc": -1, "cena_netto": 100,
                                    "netto": -100, "brutto": -123, "vat": "23%"}]),
        )

    def test_negative_correction_is_rejected_before_any_side_effect(self):
        original = copy.deepcopy(self.correction)
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "correction.xlsx"
            with patch.object(exporter.xlsxwriter, "Workbook") as workbook, \
                 patch.object(exporter, "_get_conn") as connection, \
                 patch.object(exporter, "get_setting_bool") as settings:
                with self.assertRaises(exporter.SaldeoCorrectionExportError) as caught:
                    exporter.export_saldeo_xlsx([self.correction], str(output))
                workbook.assert_not_called()
                connection.assert_not_called()
                settings.assert_not_called()
            self.assertFalse(output.exists())
        self.assertEqual(self.correction, original)
        self.assertEqual(caught.exception.corrections[0]["id"], 11)
        self.assertIn("FV/001/09/2026", str(caught.exception))
        self.assertIn("KOR/001/09/2026", str(caught.exception))
        self.assertIn("Saldeo otwórz", str(caught.exception))

    def test_mixed_batch_is_all_or_nothing_and_preserves_existing_file(self):
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "already-saved.xlsx"
            output.write_bytes(b"existing user file")
            with self.assertRaises(exporter.SaldeoCorrectionExportError):
                exporter.export_saldeo_xlsx([self.invoice, self.correction], str(output))
            self.assertEqual(output.read_bytes(), b"existing user file")

    def test_positive_and_zero_corrections_are_not_exported_as_ordinary_invoices(self):
        for net, gross in ((100, 123), (0, 0)):
            correction = dict(self.correction, kwota_netto=net, kwota_brutto=gross,
                              pozycje_json=[{"nazwa": "Zmiana", "ilosc": 1, "cena_netto": net}])
            with self.subTest(net=net), self.assertRaises(exporter.SaldeoCorrectionExportError):
                exporter.validate_saldeo_invoices([correction])

    def test_original_reference_also_identifies_correction_without_kind(self):
        correction = dict(self.invoice, korygowana_faktura_id=10)
        with self.assertRaises(exporter.SaldeoCorrectionExportError):
            exporter.validate_saldeo_invoices([correction])

    def test_negative_ordinary_line_is_rejected_even_if_total_is_positive(self):
        for key in ("ilosc", "quantity", "cena_netto", "cena_brutto", "netto", "wartosc_netto", "brutto"):
            invoice = copy.deepcopy(self.invoice)
            invoice["pozycje_json"][0][key] = -1
            with self.subTest(key=key), self.assertRaisesRegex(exporter.SaldeoExportError, "ujemną"):
                exporter.validate_saldeo_invoices([invoice])

    def test_negative_ordinary_header_is_rejected_without_positions(self):
        invoice = dict(self.invoice, pozycje_json=[], kwota_brutto=-123)
        with self.assertRaisesRegex(exporter.SaldeoExportError, "ujemną"):
            exporter.validate_saldeo_invoices([invoice])

    def test_corrupt_positions_cannot_silently_produce_empty_invoice_export(self):
        for positions in ("not json", {}, ["not a position"]):
            with self.subTest(positions=positions), self.assertRaises(exporter.SaldeoExportError):
                exporter.validate_saldeo_invoices([dict(self.invoice, pozycje_json=positions)])

    def test_positive_ordinary_invoice_still_exports_unchanged(self):
        original = copy.deepcopy(self.invoice)
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "invoice.xlsx"
            with patch.object(exporter, "_get_conn", return_value=FakeConnection()), \
                 patch.object(exporter, "get_setting_bool", return_value=False), \
                 patch.object(exporter, "get_setting_first", side_effect=lambda keys, default="": default):
                exporter.export_saldeo_xlsx([self.invoice], str(output))
            book = load_workbook(output, read_only=True, data_only=True)
            try:
                sheet = book.active
                self.assertEqual(sheet.max_row, 2)
                quantity = exporter.SALDEO_HEADERS.index("Ilość") + 1
                price = exporter.SALDEO_HEADERS.index("Cena jedn. netto") + 1
                self.assertEqual(sheet.cell(2, quantity).value, "1")
                self.assertEqual(sheet.cell(2, price).value, "100,00")
            finally:
                book.close()
        self.assertEqual(self.invoice, original)


if __name__ == "__main__":
    unittest.main()
