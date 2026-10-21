# app/reports/reports_ui.py
from app.utils.ui_imports import *
from PyQt6.QtGui import QColor
from PyQt6.QtCharts import (
    QChart, QChartView,
    QBarSeries, QBarSet,
    QLineSeries,
    QValueAxis, QBarCategoryAxis
)
from PyQt6.QtGui import QPainter
import datetime
import logging

from app.reports.reports_logic import (
    get_all_companies,
    get_months_range,
    build_report,
    build_yearly_report,
    compare_companies,
    build_profit_report,
    build_summary_report,
)
from app.reports.import_parse import SEKCJE, SEKCJA_LABELS

PROFIT_MODE = "Rentowność (import)"
PROFIT_HEADERS = ["Firma", "Miesiąc", "Ilość dokumentów", "Pracownik", "Czas pracy",
                  "Cena netto", "Koszt (stawka × czas)", "Zysk / strata"]


def format_hours(seconds: float) -> str:
    """Konwersja sekund → H:MM"""
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    return f"{h}:{m:02d}"


class ReportsPage(QWidget):
    def __init__(self, pracownik_id):
        super().__init__()
        self.pracownik_id = pracownik_id
        self._build_ui()

    def _build_ui(self):
        from PyQt6.QtWidgets import QScrollArea

        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QScrollArea.Shape.NoFrame)
        content = QWidget()
        scroll.setWidget(content)
        outer.addWidget(scroll)
        layout = QVBoxLayout(content)
        layout.setContentsMargins(12, 12, 12, 12)

        title = QLabel("<b>Raporty</b>")
        title.setFont(QFont("Arial", 18))
        layout.addWidget(title)

        mode_box = QHBoxLayout()
        self.mode_combo = QComboBox()
        self.mode_combo.addItems([
            "Raport miesięczny",
            "Raport roczny",
            "Porównanie firm",
            "Zarobek na godzinę",
            "Zestawienie",
            PROFIT_MODE,
        ])
        mode_box.addWidget(QLabel("Tryb:"))
        mode_box.addWidget(self.mode_combo)
        layout.addLayout(mode_box)

        self.firm_list = QListWidget()
        self.firm_list.setSelectionMode(QListWidget.SelectionMode.ExtendedSelection)
        self.company_names = {}
        for f in get_all_companies(self.pracownik_id):
            self.company_names[f["id"]] = f["nazwa"]
            item = QListWidgetItem(f"{f['nazwa']} (ID {f['id']})")
            item.setData(Qt.ItemDataRole.UserRole, f["id"])
            self.firm_list.addItem(item)
        layout.addWidget(self.firm_list)

        self.month_start = QComboBox()
        self.month_end = QComboBox()
        self.year_combo = QComboBox()

        now = datetime.datetime.now()
        first_year = 2026
        months = [f"{y}-{m:02d}" for y in range(first_year, max(now.year, first_year) + 1) for m in range(1, 13)]

        self.month_start.addItems(months)
        self.month_end.addItems(months)
        self.year_combo.addItems([str(y) for y in range(first_year, max(now.year, first_year) + 1)])

        self.month_box_container = QWidget()
        self.month_box = QHBoxLayout(self.month_box_container)
        self.month_box.addWidget(QLabel("Od:"))
        self.month_box.addWidget(self.month_start)
        self.month_box.addWidget(QLabel("Do:"))
        self.month_box.addWidget(self.month_end)
        layout.addWidget(self.month_box_container)

        self.year_box_container = QWidget()
        self.year_box = QHBoxLayout(self.year_box_container)
        self.year_box.addWidget(QLabel("Rok:"))
        self.year_box.addWidget(self.year_combo)
        layout.addWidget(self.year_box_container)

        self.mode_combo.currentTextChanged.connect(self._update_mode_ui)

        btn = QPushButton("Generuj raport")
        btn.setProperty("class", "primary")
        btn.clicked.connect(self._generate_report)
        layout.addWidget(btn)

        btn_export = QPushButton("Eksport do Excel")
        btn_export.clicked.connect(self._export_to_excel)
        layout.addWidget(btn_export)

        import_box = QHBoxLayout()
        btn_import_work = QPushButton("Importuj czas pracy (Excel)")
        btn_import_work.clicked.connect(self._import_work)
        btn_import_inv = QPushButton("Importuj faktury (Excel)")
        btn_import_inv.clicked.connect(self._import_invoices)
        btn_rates = QPushButton("Stawki pracowników")
        btn_rates.clicked.connect(self._edit_rates)
        btn_delete = QPushButton("Usuń zaimportowane dane")
        btn_delete.clicked.connect(self._delete_imported)
        for b in (btn_import_work, btn_import_inv, btn_rates, btn_delete):
            import_box.addWidget(b)
        layout.addLayout(import_box)

        self.import_label = QLabel("<b>Księgowość</b>")
        layout.addWidget(self.import_label)
        self.import_table = self._make_profit_table()
        layout.addWidget(self.import_table)
        self.import_label_kadry = QLabel("<b>Kadry</b>")
        layout.addWidget(self.import_label_kadry)
        self.import_table_kadry = self._make_profit_table()
        layout.addWidget(self.import_table_kadry)
        self.chart_top = QChartView()
        self.chart_top.setRenderHint(QPainter.RenderHint.Antialiasing)

        self.chart_bottom = QChartView()
        self.chart_bottom.setRenderHint(QPainter.RenderHint.Antialiasing)

        layout.addWidget(self.chart_top)
        layout.addWidget(self.chart_bottom)
        for chart_view in (self.chart_top, self.chart_bottom):
            chart_view.setMinimumHeight(320)
        self.firm_list.setMinimumHeight(140)

        self.summary_table = QTableWidget(0, 6)
        self.summary_table.setHorizontalHeaderLabels(
            ["Klient", "Księgowa", "Miesiąc", "Typ", "Czas pracy", "Cena za usługę"]
        )
        self.summary_table.horizontalHeader().setStretchLastSection(True)
        self.summary_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        layout.addWidget(self.summary_table)

        self.report_status = QLabel("Wybierz firmy i zakres, aby wygenerować raport.")
        self.report_status.setStyleSheet(
            "QLabel { background: #f4f6f8; border: 1px solid #d8dde3; "
            "border-radius: 6px; padding: 8px; color: #263238; }"
        )
        layout.addWidget(self.report_status)

        self.report_table = QTableWidget(0, 6)
        self.report_table.setHorizontalHeaderLabels(
            ["Klient", "Miesiąc", "Księgowość", "Kadry", "Czas księgowości", "Czas kadr"]
        )
        self.report_table.setAlternatingRowColors(True)
        self.report_table.setSortingEnabled(True)
        self.report_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.report_table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(self.report_table)

        self._update_mode_ui(self.mode_combo.currentText())
        for table in (self.summary_table, self.report_table, self.import_table, self.import_table_kadry):
            table.setMinimumHeight(300)

    @staticmethod
    def _make_profit_table():
        table = QTableWidget(0, len(PROFIT_HEADERS))
        table.setHorizontalHeaderLabels(PROFIT_HEADERS)
        header = table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        header.setStretchLastSection(True)
        table.setAlternatingRowColors(True)
        table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        return table

    def _update_mode_ui(self, mode):
        profit = mode == PROFIT_MODE
        summary = mode == "Zestawienie"
        self.month_box_container.setVisible(
            mode in ["Raport miesięczny", "Porównanie firm", "Zarobek na godzinę", "Zestawienie", PROFIT_MODE]
        )
        self.year_box_container.setVisible(mode == "Raport roczny")
        self.summary_table.setVisible(summary)
        self.import_table.setVisible(profit)
        self.import_label.setVisible(profit)
        self.import_table_kadry.setVisible(profit)
        self.import_label_kadry.setVisible(profit)
        self.report_table.setVisible(not summary and not profit)
        self.chart_top.setVisible(not summary and not profit)
        self.chart_bottom.setVisible(not summary and not profit)

    def _generate_report(self):
        selected = [
            self.firm_list.item(i).data(Qt.ItemDataRole.UserRole)
            for i in range(self.firm_list.count())
            if self.firm_list.item(i).isSelected()
        ]

        if not selected:
            QMessageBox.warning(self, "Brak wyboru", "Wybierz co najmniej jedną firmę.")
            return

        mode = self.mode_combo.currentText()

        if mode == "Raport miesięczny":
            start = self.month_start.currentText()
            end = self.month_end.currentText()
            months = get_months_range(start, end)
            data = build_report(selected, months, self.pracownik_id)
            self._render_mixed(data, selected, months)

        elif mode == "Raport roczny":
            year = int(self.year_combo.currentText())
            if len(selected) != 1:
                QMessageBox.warning(self, "Błąd", "Raport roczny wymaga jednej firmy.")
                return
            firma_id = selected[0]
            months = [f"{year}-{m:02d}" for m in range(1, 13)]
            data = build_yearly_report(firma_id, year)
            self._render_mixed({firma_id: data}, [firma_id], months)

        elif mode == "Porównanie firm":
            miesiac = self.month_start.currentText()
            data = compare_companies(selected, miesiac, self.pracownik_id)
            self._render_compare(data, selected, miesiac)

        elif mode == "Zarobek na godzinę":
            start = self.month_start.currentText()
            end = self.month_end.currentText()
            months = get_months_range(start, end)
            data = build_profit_report(selected, months, self.pracownik_id)
            self._render_profit(data, selected, months)

        elif mode == "Zestawienie":
            start = self.month_start.currentText()
            end = self.month_end.currentText()
            months = get_months_range(start, end)
            self._render_summary(build_summary_report(selected, months, self.pracownik_id))

        elif mode == PROFIT_MODE:
            months = get_months_range(self.month_start.currentText(), self.month_end.currentText())
            self._render_import_report(self._profit_rows(selected, months))

    def _profit_rows(self, selected, months):
        from app.reports.import_dao import build_profitability_report

        return build_profitability_report(selected, months, self.pracownik_id)

    @staticmethod
    def _profit_values(item, raw=False):
        docs = "" if item["dokumenty"] is None else item["dokumenty"]
        if raw:
            return [item["klient"], item["miesiac"], docs, item["pracownik"],
                    format_hours(item["sekundy"]), item["cena"], item["koszt"], item["zysk"]]
        return [item["klient"], item["miesiac"], docs, item["pracownik"],
                format_hours(item["sekundy"]), f"{item['cena']:.2f} zł",
                f"{item['koszt']:.2f} zł", f"{item['zysk']:.2f} zł"]

    def _render_import_report(self, rows):
        tables = {"ksiegowosc": self.import_table, "kadry": self.import_table_kadry}
        totals = {}
        for sekcja, table in tables.items():
            table.setRowCount(0)
            sek = cena = koszt = 0.0
            for item in rows:
                if item["sekcja"] != sekcja:
                    continue
                row = table.rowCount()
                table.insertRow(row)
                for col, value in enumerate(self._profit_values(item)):
                    cell = QTableWidgetItem(str(value))
                    if col >= 4:
                        cell.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                    if col == 7 and item["zysk"] < 0:
                        cell.setForeground(QColor("#c62828"))
                    table.setItem(row, col, cell)
                sek += item["sekundy"]
                cena += item["cena"]
                koszt += item["koszt"]
            totals[sekcja] = (sek, cena, koszt)
        if not rows:
            self.report_status.setText("Brak zaimportowanych danych dla wybranych firm i zakresu.")
            return
        self.report_status.setText(" | ".join(
            f"{SEKCJA_LABELS[s]}: czas {format_hours(t[0])}, cena {t[1]:.2f} zł, "
            f"koszt {t[2]:.2f} zł, zysk {t[1] - t[2]:.2f} zł"
            for s, t in totals.items()
        ))

    def _import_work(self):
        self._run_import("czas pracy")

    def _import_invoices(self):
        self._run_import("faktury")

    def _run_import(self, kind):
        from PyQt6.QtWidgets import QFileDialog, QInputDialog
        from app.reports import import_dao

        if not self._can_import():
            QMessageBox.warning(self, "Import", "Brak uprawnień do importu danych.")
            return
        path, _ = QFileDialog.getOpenFileName(
            self, f"Wybierz plik Excel ({kind})", "", "Pliki Excel (*.xlsx)"
        )
        if not path:
            return
        try:
            if kind == "faktury":
                options = [self.month_start.itemText(i) for i in range(self.month_start.count())]
                month, ok = QInputDialog.getItem(
                    self, "Miesiąc faktur",
                    "Miesiąc dla faktur, w których numerze nie ma miesiąca (np. FV/0001/01/2026):",
                    options, 0, False,
                )
                if not ok:
                    return
                result = import_dao.import_invoices_file(path, month)
            else:
                result = import_dao.import_work_file(path)
        except Exception as exc:
            logging.getLogger(__name__).exception("Import raportów nie powiódł się")
            QMessageBox.critical(self, "Import", f"Nie udało się zaimportować pliku:\n{exc}")
            return
        QMessageBox.information(
            self, "Import",
            f"Wczytano wierszy (od 2026): {result['wiersze']}\n"
            f"Dodano: {result['dodane']}\n"
            f"Pominięto (już zaimportowane): {result['pominiete']}\n"
            f"Bez dopasowanej firmy: {result['bez_firmy']}",
        )

    def _delete_imported(self):
        from PyQt6.QtWidgets import QInputDialog
        from app.reports import import_dao

        if not self._can_import():
            QMessageBox.warning(self, "Usuwanie", "Brak uprawnień do usuwania danych.")
            return
        labels = {"Czas pracy": "czas", "Faktury": "faktury", "Czas pracy i faktury": "oba"}
        choice, ok = QInputDialog.getItem(
            self, "Usuń zaimportowane dane", "Co usunąć?", list(labels), 0, False
        )
        if not ok:
            return
        months = get_months_range(self.month_start.currentText(), self.month_end.currentText())
        selected = [
            self.firm_list.item(i).data(Qt.ItemDataRole.UserRole)
            for i in range(self.firm_list.count())
            if self.firm_list.item(i).isSelected()
        ]
        scope = f"{len(selected)} zaznaczonych firm" if selected else "WSZYSTKICH firm (łącznie z niedopasowanymi)"
        if QMessageBox.question(
            self, "Potwierdź usunięcie",
            f"Usunąć „{choice}” z zakresu {months[0]} – {months[-1]} dla {scope}?\n"
            "Tej operacji nie można cofnąć.",
        ) != QMessageBox.StandardButton.Yes:
            return
        try:
            deleted = import_dao.delete_imported(labels[choice], months, selected or None)
        except Exception as exc:
            logging.getLogger(__name__).exception("Usuwanie zaimportowanych danych nie powiodło się")
            QMessageBox.critical(self, "Usuwanie", f"Nie udało się usunąć danych:\n{exc}")
            return
        QMessageBox.information(self, "Usuwanie", f"Usunięto wierszy: {deleted}")

    def _can_import(self):
        from app.security.authorization import can, is_administrator

        return is_administrator(self.pracownik_id) or can(self.pracownik_id, "can_view_finances")

    def _edit_rates(self):
        from app.reports import import_dao

        if not self._can_import():
            QMessageBox.warning(self, "Stawki", "Brak uprawnień do edycji stawek.")
            return
        employees = import_dao.list_employee_rates()
        dlg = QDialog(self)
        dlg.setWindowTitle("Stawki godzinowe pracowników (koszt pracy, zł/h)")
        lay = QVBoxLayout(dlg)
        table = QTableWidget(len(employees), 2)
        table.setHorizontalHeaderLabels(["Pracownik", "Stawka zł/h (uśredniona)"])
        table.horizontalHeader().setStretchLastSection(True)
        for i, emp in enumerate(employees):
            name = QTableWidgetItem(emp["nazwa"])
            name.setFlags(name.flags() & ~Qt.ItemFlag.ItemIsEditable)
            name.setData(Qt.ItemDataRole.UserRole, emp["id"])
            table.setItem(i, 0, name)
            table.setItem(i, 1, QTableWidgetItem("" if emp["stawka"] is None else f"{emp['stawka']:.2f}"))
        lay.addWidget(table)
        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Save | QDialogButtonBox.StandardButton.Cancel
        )
        buttons.accepted.connect(dlg.accept)
        buttons.rejected.connect(dlg.reject)
        lay.addWidget(buttons)
        if not dlg.exec():
            return
        rates = {}
        for i in range(table.rowCount()):
            raw = (table.item(i, 1).text() or "").strip().replace(",", ".")
            try:
                value = float(raw) if raw else None
            except ValueError:
                QMessageBox.warning(self, "Stawki", f"Niepoprawna stawka: {raw}")
                return
            rates[table.item(i, 0).data(Qt.ItemDataRole.UserRole)] = value
        import_dao.set_employee_rates(rates)

    def _render_summary(self, rows):
        self.summary_table.setRowCount(0)
        for item in rows:
            row = self.summary_table.rowCount()
            self.summary_table.insertRow(row)
            values = [
                item["klient"],
                item["ksiegowa"],
                item["miesiac"],
                "Kadry" if item["typ"] == "kadry" else "Księgowość",
                format_hours(item["sekundy"]),
                f"{item['cena_za_usluge']:.2f} zł",
            ]
            for col, value in enumerate(values):
                self.summary_table.setItem(row, col, QTableWidgetItem(str(value)))

    # ============================================================
    # WYKRES MIESZANY — SŁUPKI (KWOTA) + LINIA (CZAS)
    # ============================================================

    def _render_mixed(self, data, selected, months):
        self._render_report_table(data, selected, months)
        categories = months

        # --- Kadry ---
        chart_kadry = QChart()
        chart_kadry.setTitle("Kadry — kwota vs czas")

        bar_series = QBarSeries()
        line_series = QLineSeries()
        line_series.setName("Czas (godziny)")

        # słupki
        for fid in selected:
            bar = QBarSet(f"Firma {fid}")
            for m in months:
                bar.append(data[fid][m]["koszt_kadry"])
            bar_series.append(bar)

        # linia — poprawione indeksowanie
        for fid in selected:
            for idx, m in enumerate(months):
                seconds = data[fid][m]["godziny_kadry"]
                hours_float = seconds / 3600
                line_series.append(idx, hours_float)
                line_series.setPointLabelsVisible(True)
                line_series.setPointLabelsFormat(format_hours(seconds))

        chart_kadry.addSeries(bar_series)
        chart_kadry.addSeries(line_series)

        axis_x = QBarCategoryAxis()
        axis_x.append(categories)
        chart_kadry.addAxis(axis_x, Qt.AlignmentFlag.AlignBottom)
        bar_series.attachAxis(axis_x)
        line_series.attachAxis(axis_x)

        axis_y = QValueAxis()
        axis_y.setTitleText("Kwota (zł)")
        chart_kadry.addAxis(axis_y, Qt.AlignmentFlag.AlignLeft)
        bar_series.attachAxis(axis_y)

        axis_y2 = QValueAxis()
        axis_y2.setTitleText("Czas (godziny)")
        chart_kadry.addAxis(axis_y2, Qt.AlignmentFlag.AlignRight)
        line_series.attachAxis(axis_y2)

        self.chart_top.setChart(chart_kadry)

        # --- Księgowość ---
        chart_ks = QChart()
        chart_ks.setTitle("Księgowość — kwota vs czas")

        bar_series2 = QBarSeries()
        line_series2 = QLineSeries()
        line_series2.setName("Czas (godziny)")

        for fid in selected:
            bar = QBarSet(f"Firma {fid}")
            for m in months:
                bar.append(data[fid][m]["koszt_ksiegowosc"])
            bar_series2.append(bar)

        for fid in selected:
            for idx, m in enumerate(months):
                seconds = data[fid][m]["godziny_ksiegowosc"]
                hours_float = seconds / 3600
                line_series2.append(idx, hours_float)
                line_series2.setPointLabelsVisible(True)
                line_series2.setPointLabelsFormat(format_hours(seconds))

        chart_ks.addSeries(bar_series2)
        chart_ks.addSeries(line_series2)

        axis_x2 = QBarCategoryAxis()
        axis_x2.append(categories)
        chart_ks.addAxis(axis_x2, Qt.AlignmentFlag.AlignBottom)
        bar_series2.attachAxis(axis_x2)
        line_series2.attachAxis(axis_x2)

        axis_y3 = QValueAxis()
        axis_y3.setTitleText("Kwota (zł)")
        chart_ks.addAxis(axis_y3, Qt.AlignmentFlag.AlignLeft)
        bar_series2.attachAxis(axis_y3)

        axis_y4 = QValueAxis()
        axis_y4.setTitleText("Czas (godziny)")
        chart_ks.addAxis(axis_y4, Qt.AlignmentFlag.AlignRight)
        line_series2.attachAxis(axis_y4)

        self.chart_bottom.setChart(chart_ks)

    def _render_report_table(self, data, selected, months):
        self.report_table.setSortingEnabled(False)
        self.report_table.setRowCount(0)
        total_ks = total_kadry = total_ks_time = total_kadry_time = 0.0
        for fid in selected:
            for month in months:
                row_data = data[fid][month]
                row = self.report_table.rowCount()
                self.report_table.insertRow(row)
                values = [
                    self.company_names.get(fid, f"Firma {fid}"),
                    month,
                    f"{row_data.get('koszt_ksiegowosc', 0):.2f} zł",
                    f"{row_data.get('koszt_kadry', 0):.2f} zł",
                    format_hours(row_data.get("godziny_ksiegowosc", 0)),
                    format_hours(row_data.get("godziny_kadry", 0)),
                ]
                for column, value in enumerate(values):
                    self.report_table.setItem(row, column, QTableWidgetItem(str(value)))
                total_ks += float(row_data.get("koszt_ksiegowosc", 0) or 0)
                total_kadry += float(row_data.get("koszt_kadry", 0) or 0)
                total_ks_time += float(row_data.get("godziny_ksiegowosc", 0) or 0)
                total_kadry_time += float(row_data.get("godziny_kadry", 0) or 0)
        self.report_table.setSortingEnabled(True)
        self.report_status.setText(
            f"Podsumowanie: księgowość {total_ks:.2f} zł / {format_hours(total_ks_time)}, "
            f"kadry {total_kadry:.2f} zł / {format_hours(total_kadry_time)}"
        )

    # ============================================================
    # ZAROBEK NA GODZINĘ — SŁUPKI
    # ============================================================

    def _render_profit(self, data, selected, months):
        categories = months

        chart_kadry = QChart()
        chart_kadry.setTitle("Zarobek na godzinę — Kadry")

        series_kadry = QBarSeries()

        for fid in selected:
            bar = QBarSet(f"Firma {fid}")
            for m in months:
                bar.append(data[fid][m]["profit_kadry"])
            series_kadry.append(bar)

        chart_kadry.addSeries(series_kadry)

        axis_x = QBarCategoryAxis()
        axis_x.append(categories)
        chart_kadry.addAxis(axis_x, Qt.AlignmentFlag.AlignBottom)
        series_kadry.attachAxis(axis_x)

        axis_y = QValueAxis()
        axis_y.setTitleText("Zł / godz")
        chart_kadry.addAxis(axis_y, Qt.AlignmentFlag.AlignLeft)
        series_kadry.attachAxis(axis_y)

        self.chart_top.setChart(chart_kadry)

        chart_ks = QChart()
        chart_ks.setTitle("Zarobek na godzinę — Księgowość")

        series_ks = QBarSeries()

        for fid in selected:
            bar = QBarSet(f"Firma {fid}")
            for m in months:
                bar.append(data[fid][m]["profit_ksiegowosc"])
            series_ks.append(bar)

        chart_ks.addSeries(series_ks)

        axis_x2 = QBarCategoryAxis()
        axis_x2.append(categories)
        chart_ks.addAxis(axis_x2, Qt.AlignmentFlag.AlignBottom)
        series_ks.attachAxis(axis_x2)

        axis_y2 = QValueAxis()
        axis_y2.setTitleText("Zł / godz")
        chart_ks.addAxis(axis_y2, Qt.AlignmentFlag.AlignLeft)
        series_ks.attachAxis(axis_y2)

        self.chart_bottom.setChart(chart_ks)

    # ============================================================
    # PORÓWNANIE FIRM — LINIOWY
    # ============================================================

    def _render_compare(self, data, selected, miesiac):
        chart_top = QChart()
        chart_top.setTitle(f"Kadry — porównanie firm ({miesiac})")

        chart_bottom = QChart()
        chart_bottom.setTitle(f"Księgowość — porównanie firm ({miesiac})")

        s_kadry = QLineSeries()
        s_ksieg = QLineSeries()

        x = 1
        for firma_id in selected:
            s_kadry.append(x, data[firma_id]["koszt_na_godzine_kadry"])
            s_ksieg.append(x, data[firma_id]["koszt_na_godzine_ksiegowosc"])
            x += 1

        chart_top.addSeries(s_kadry)
        chart_bottom.addSeries(s_ksieg)

        chart_top.createDefaultAxes()
        chart_bottom.createDefaultAxes()

        self.chart_top.setChart(chart_top)
        self.chart_bottom.setChart(chart_bottom)

    # ============================================================
    # EKSPORT DO EXCELA
    # ============================================================

    def _export_to_excel(self):
        import openpyxl
        from openpyxl.styles import Font
        from PyQt6.QtWidgets import QFileDialog
        from app.utils.xlsx_safe import safe_row

        mode = self.mode_combo.currentText()

        selected = [
            self.firm_list.item(i).data(Qt.ItemDataRole.UserRole)
            for i in range(self.firm_list.count())
            if self.firm_list.item(i).isSelected()
        ]

        if not selected:
            QMessageBox.warning(self, "Błąd", "Wybierz firmę przed eksportem.")
            return

        if mode == "Raport roczny":
            year = int(self.year_combo.currentText())
            months = [f"{year}-{m:02d}" for m in range(1, 13)]
            data = build_yearly_report(selected[0], year)
            data = {selected[0]: data}
        else:
            start = self.month_start.currentText()
            end = self.month_end.currentText()
            months = get_months_range(start, end)

            if mode == "Raport miesięczny":
                data = build_report(selected, months, self.pracownik_id)
            elif mode == "Zestawienie":
                rows = build_summary_report(selected, months, self.pracownik_id)
                filepath, _ = QFileDialog.getSaveFileName(
                    self, "Zapisz zestawienie Excel", "", "Pliki Excel (*.xlsx)"
                )
                if not filepath:
                    return
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Zestawienie"
                ws.append(["Klient", "Księgowa", "Miesiąc", "Typ", "Czas pracy", "Cena za usługę"])
                for row in rows:
                    ws.append(safe_row([
                        row["klient"], row["ksiegowa"], row["miesiac"],
                        row["typ"], format_hours(row["sekundy"]),
                        row["cena_za_usluge"],
                    ]))
                wb.save(filepath)
                QMessageBox.information(self, "Sukces", f"Zestawienie zapisane:\n{filepath}")
                return
            elif mode == PROFIT_MODE:
                rows = self._profit_rows(selected, months)
                filepath, _ = QFileDialog.getSaveFileName(
                    self, "Zapisz rentowność Excel", "", "Pliki Excel (*.xlsx)"
                )
                if not filepath:
                    return
                wb = openpyxl.Workbook()
                wb.remove(wb.active)
                for sekcja in SEKCJE:
                    ws = wb.create_sheet(SEKCJA_LABELS[sekcja])
                    ws.append(PROFIT_HEADERS)
                    for item in rows:
                        if item["sekcja"] != sekcja:
                            continue
                        ws.append(safe_row(self._profit_values(item, raw=True)))
                wb.save(filepath)
                QMessageBox.information(self, "Sukces", f"Raport zapisany:\n{filepath}")
                return
            elif mode == "Zarobek na godzinę":
                data = build_profit_report(selected, months, self.pracownik_id)
            elif mode == "Porównanie firm":
                data = compare_companies(
                    selected,
                    self.month_start.currentText(),
                    self.pracownik_id,
                )
                months = [self.month_start.currentText()]

        filepath, _ = QFileDialog.getSaveFileName(
            self,
            "Zapisz raport Excel",
            "",
            "Pliki Excel (*.xlsx)"
        )
        if not filepath:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Raport"

        headers = [
            "Firma ID",
            "Miesiąc",
            "Kwota Kadry",
            "Kwota Księgowość",
            "Czas Kadry (H:MM)",
            "Czas Księgowość (H:MM)",
            "Zarobek/h Kadry",
            "Zarobek/h Księgowość"
        ]
        ws.append(headers)

        for col in range(1, len(headers) + 1):
            ws.cell(row=1, column=col).font = Font(bold=True)

        for fid in selected:
            for m in months:
                row = data[fid][m]

                godz_kadry = format_hours(row.get("godziny_kadry", 0))
                godz_ks = format_hours(row.get("godziny_ksiegowosc", 0))

                ws.append([
                    fid,
                    m,
                    row.get("koszt_kadry", row.get("profit_kadry", 0)),
                    row.get("koszt_ksiegowosc", row.get("profit_ksiegowosc", 0)),
                    godz_kadry,
                    godz_ks,
                    round(row.get("profit_kadry", row.get("koszt_na_godzine_kadry", 0)), 2),
                    round(row.get("profit_ksiegowosc", row.get("koszt_na_godzine_ksiegowosc", 0)), 2)
                ])

        try:
            wb.save(filepath)
            QMessageBox.information(self, "Sukces", f"Raport zapisany:\n{filepath}")
        except Exception as e:
            QMessageBox.critical(self, "Błąd", f"Nie udało się zapisać pliku:\n{e}")
