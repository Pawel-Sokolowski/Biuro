"""Regresja: nazwa Saldeo i globalne konto mają pierwszeństwo w XLSX."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from app.utils import saldeo_exporter


class FakeCursor:
    def execute(self, sql, params):
        pass

    def fetchone(self):
        return {
            "nazwa": "Nazwa pełna kontrahenta",
            "nazwa_skrocona": "Nazwa wewnętrzna",
            "nazwa_saldeo": "Nazwa Saldeo",
            "nip": "1234563218",
            "regon": "",
            "miasto": "Szczecin",
            "kod_pocztowy": "70-000",
            "ulica": "",
            "adres_budynek": "",
            "adres_lokal": "",
            "adres_email": "kontakt@example.com",
            "metoda_kasowa": None,
            "wysylac_do_ksef": 0,
            "wysylac_do_kontrahenta": 0,
        }


class FakeConnection:
    def cursor(self):
        return FakeCursor()

    def close(self):
        pass


class SaldeoExportNameTests(unittest.TestCase):
    def test_export_uses_saldeo_short_name_and_global_bank_account(self):
        invoice = {
            "firma_id": 1, "miesiac": "2026-10",
            "nabywca_skrocona": "Stara nazwa na fakturze",
            "konto_bankowe": "99999999999999999999999999",
            "pozycje_json": [{"nazwa": "Usługa", "ilosc": 1,
                              "cena_netto": 100, "vat": "23%"}],
        }

        def setting(keys, default=""):
            if "konto_bankowe_biura" in keys:
                return "11111111111111111111111111"
            return default

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "faktury.xlsx"
            with patch.object(saldeo_exporter, "_get_conn", return_value=FakeConnection()), \
                 patch.object(saldeo_exporter, "get_setting_bool", return_value=False), \
                 patch.object(saldeo_exporter, "get_setting_first", side_effect=setting):
                saldeo_exporter.export_saldeo_xlsx([invoice], str(output))

            book = load_workbook(output, read_only=True, data_only=True)
            try:
                sheet = book.active
                name_col = saldeo_exporter.SALDEO_HEADERS.index("Nabywca (nazwa skrócona kontrahenta)") + 1
                bank_col = saldeo_exporter.SALDEO_HEADERS.index("Konto bankowe") + 1
                self.assertEqual(sheet.cell(2, name_col).value, "Nazwa Saldeo")
                self.assertEqual(sheet.cell(2, bank_col).value, "11111111111111111111111111")
            finally:
                book.close()


if __name__ == "__main__":
    unittest.main()
