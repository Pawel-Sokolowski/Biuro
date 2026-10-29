"""Import nazw skróconych kontrahentów z Saldeo bez zmiany nazw rejestrowych."""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from app.db.dao_conn import get_conn
from app.security.authorization import is_administrator


def _header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").replace("\ufeff", ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]", "", text.lower())


def _cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _nip(value: Any) -> str:
    raw = _cell(value)
    # Excel może zapisać liczbę jako 1234567890.0.
    if re.fullmatch(r"\d{10}\.0+", raw):
        raw = raw.split(".", 1)[0]
    digits = re.sub(r"\D", "", raw)
    return digits if len(digits) == 10 else ""


def _csv_rows(path: Path) -> list[list[Any]]:
    data = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "utf-16", "cp1250"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeError:
            pass
    if text is None:
        raise ValueError("Nie udało się odczytać kodowania pliku CSV.")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        return list(csv.reader(io.StringIO(text), dialect=dialect))
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _xlsx_rows(path: Path) -> list[list[Any]]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        book.close()


def read_saldeo_contractors(path: str) -> list[dict[str, str]]:
    """Wczytuje kolumny Typ, Nazwa skrócona, NIP i Adres email."""
    source = Path(path)
    extension = source.suffix.lower()
    if extension == ".csv":
        rows = _csv_rows(source)
    elif extension == ".xlsx":
        rows = _xlsx_rows(source)
    else:
        raise ValueError("Obsługiwane formaty to CSV i XLSX.")

    required = {"typ", "nazwaskrocona", "nip"}
    aliases = {"adresemail": "email", "email": "email", "nazwaskrocona": "short_name"}
    columns = None
    header_index = None
    for i, row in enumerate(rows[:20]):
        normalized = [_header(cell) for cell in row]
        if required.issubset(set(normalized)):
            columns = {aliases.get(name, name): j for j, name in enumerate(normalized)}
            header_index = i
            break
    if columns is None:
        raise ValueError("Brak nagłówków Saldeo: Typ, Nazwa skrócona, NIP.")

    records = []
    for row_number, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not any(_cell(v) for v in row):
            continue

        def value(name: str) -> str:
            index = columns.get(name)
            return _cell(row[index]) if index is not None and index < len(row) else ""

        records.append({
            "row": str(row_number),
            "typ": value("typ"),
            "nip": value("nip"),
            "short_name": value("short_name"),
            "email": value("email"),
        })
    return records


def import_saldeo_contractors(path: str, actor_id: int) -> dict[str, Any]:
    """Dopasowuje wyłącznie istniejące firmy po NIP; wymaga roli administratora.

    Zmienia tylko firmy.nazwa_saldeo oraz puste firmy.adres_email.
    """
    if not is_administrator(actor_id):
        raise PermissionError("Import danych Saldeo jest dostępny tylko dla administratora.")

    records = read_saldeo_contractors(path)
    result: dict[str, Any] = {
        "total": len(records), "eligible": 0, "skipped_type": 0,
        "invalid_nip": 0, "missing_company": 0, "ambiguous": 0,
        "duplicates": 0, "updated_names": 0, "filled_emails": 0,
        "unchanged": 0, "warnings": [],
    }

    by_nip: dict[str, dict[str, str]] = {}
    conflicts: set[str] = set()
    for entry in records:
        if re.sub(r"\s+", "", entry["typ"]).upper() != "K/D":
            result["skipped_type"] += 1
            continue
        result["eligible"] += 1
        nip = _nip(entry["nip"])
        if not nip:
            result["invalid_nip"] += 1
            result["warnings"].append("Wiersz {}: nieprawidłowy NIP.".format(entry["row"]))
            continue
        normalized = {"short_name": entry["short_name"], "email": entry["email"]}
        if nip in by_nip:
            result["duplicates"] += 1
            if by_nip[nip] != normalized:
                conflicts.add(nip)
            continue
        by_nip[nip] = normalized

    conn = get_conn()
    try:
        rows = conn.execute("SELECT id, nip, nazwa_saldeo, adres_email FROM firmy").fetchall()
        db_by_nip: dict[str, list[dict[str, Any]]] = {}
        for row in rows:
            company = dict(row)
            nip = _nip(company["nip"])
            if nip:
                db_by_nip.setdefault(nip, []).append(company)

        for nip, incoming in by_nip.items():
            if nip in conflicts:
                result["ambiguous"] += 1
                result["warnings"].append("NIP {}: różne dane w pliku — pominięto.".format(nip))
                continue
            candidates = db_by_nip.get(nip, [])
            if not candidates:
                result["missing_company"] += 1
                continue
            if len(candidates) != 1:
                result["ambiguous"] += 1
                result["warnings"].append("NIP {}: kilka firm w bazie — pominięto.".format(nip))
                continue

            company = candidates[0]
            updates = {}
            name = incoming["short_name"]
            if name and name != (company["nazwa_saldeo"] or "").strip():
                updates["nazwa_saldeo"] = name

            email = incoming["email"]
            if email and not (company["adres_email"] or "").strip():
                if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
                    updates["adres_email"] = email
                else:
                    result["warnings"].append("NIP {}: nieprawidłowy e-mail — pominięto adres.".format(nip))
            if not updates:
                result["unchanged"] += 1
                continue

            # Ograniczamy aktualizację do dwóch dozwolonych pól.
            if "nazwa_saldeo" in updates and "adres_email" in updates:
                conn.execute(
                    "UPDATE firmy SET nazwa_saldeo = %s, adres_email = %s "
                    "WHERE id = %s AND NULLIF(BTRIM(COALESCE(adres_email, '')), '') IS NULL",
                    (updates["nazwa_saldeo"], updates["adres_email"], company["id"]),
                )
            elif "nazwa_saldeo" in updates:
                conn.execute("UPDATE firmy SET nazwa_saldeo = %s WHERE id = %s",
                             (updates["nazwa_saldeo"], company["id"]))
            else:
                conn.execute(
                    "UPDATE firmy SET adres_email = %s WHERE id = %s "
                    "AND NULLIF(BTRIM(COALESCE(adres_email, '')), '') IS NULL",
                    (updates["adres_email"], company["id"]),
                )
            if "nazwa_saldeo" in updates:
                result["updated_names"] += 1
            if "adres_email" in updates:
                result["filled_emails"] += 1
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return result
