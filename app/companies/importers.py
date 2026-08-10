import openpyxl
import sqlite3
import time
from app.utils.parsing import normalize_identifier, resolve_employee_ids

def import_companies_from_excel_file(filepath, db_path):
    workbook = openpyxl.load_workbook(filepath, data_only=True)
    worksheet = workbook.active
    header_cells = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_cells:
        return {"created": 0, "updated": 0, "warnings": ["Plik nie zawiera nagłówków."]}

    headers = [str(header).strip() if header is not None else "" for header in header_cells]
    header_index = {header: index for index, header in enumerate(headers)}
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    created_count = 0
    updated_count = 0
    warnings = []

    try:
        for row_number, row_values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value not in (None, "") for value in row_values):
                continue

            payload = {
                "uop": 0,
                "uz": 0,
                "uod": 0,
                "glowny_cennik_manualny": 0
            }
            mapping = {
                "NIP": "nip", "REGON": "regon", "KRS": "krs", "PESEL": "pesel",
                "Nazwa Firmy": "nazwa", "Nazwa skrócona": "nazwa_skrocona",
                "Forma działalności": "forma_dzialalnosci", "Forma opodatkowania": "forma_opodatkowania",
                "Analityka FK": "analityka_fk", "Osoba kontaktowa": "osoba_kontaktowa",
                "Adres email": "adres_email", "Nr telefonu": "nr_telefonu",
                "Miasto": "miasto", "Kod pocztowy": "kod_pocztowy", "Ulica": "ulica",
                "Adres zamieszkania - miasto": "adres_zam_miasto", "Adres zamieszkania - kod pocztowy": "adres_zam_kod_pocztowy",
                "Adres zamieszkania - ulica": "adres_zam_ulica", "Adres korespondencyjny - miasto": "adres_koresp_miasto",
                "Adres korespondencyjny - kod pocztowy": "adres_koresp_kod_pocztowy", "Adres korespondencyjny - ulica": "adres_koresp_ulica",
                "Data rozpoczęcia działalności gospodarczej": "data_rozpoczecia_dg",
                "Data rozpoczęcia sprzedaży dla osoby fizycznej": "data_rozpoczecia_sprzedazy",
                "Data zakończenia ulgi ZUS na start": "data_koniec_zus_start",
                "Data zakończenia ZUS preferencyjny": "data_koniec_zus_pref",
                "Data zakończenia małego ZUS PLUS": "data_koniec_maly_zus",
                "Data zakończenia A1": "data_koniec_a1", "Data rozpoczęcia obsługi": "data_rozpoczecia_obslugi",
                "Data zakończenia obsługi": "data_zakonczenia_obslugi", "Zawieszenia działalności": "zawieszenia_dzialalnosci",
                "Kody PKD": "kody_pkd", "Cenniki usług księgowych": "cennik_ksiegowy",
                "Cenniki usług kadrowych": "cennik_kadrowy", "Cenniki czasowe": "cennik_czasowy",
                "Liczba pracowników": "liczba_pracownikow", "Limit dokumentów": "limit_dokumentow"
            }
            for source_header, target_field in mapping.items():
                if source_header in header_index:
                    payload[target_field] = row_values[header_index[source_header]]

            assigned_raw = row_values[header_index["Przypisani pracownicy"]] if "Przypisani pracownicy" in header_index else None
            payload["assigned_employee_ids"], unresolved = resolve_employee_ids(cursor, assigned_raw)
            if unresolved:
                warnings.append(f"Wiersz {row_number}: nie znaleziono pracowników {', '.join(unresolved)}.")

            payload["glowny_cennik"] = payload.get("cennik_ksiegowy", "")
            company_name = str(payload.get("nazwa", "")).strip()
            if not company_name:
                warnings.append(f"Wiersz {row_number}: pominięto pustą nazwę firmy.")
                continue

            normalized_nip = normalize_identifier(payload.get("nip"), {10})
            company_id = None
            if normalized_nip:
                cursor.execute("SELECT id FROM firmy WHERE nip = ?", (normalized_nip,))
                existing_row = cursor.fetchone()
                company_id = existing_row[0] if existing_row else None

            if company_id:
                cursor.execute("UPDATE firmy SET nazwa = ?, glowny_cennik = ? WHERE id = ?", (company_name, payload["glowny_cennik"], company_id))
                updated_count += 1
            else:
                cursor.execute("INSERT INTO firmy (nazwa, nip, glowny_cennik) VALUES (?, ?, ?)", (company_name, normalized_nip, payload["glowny_cennik"]))
                created_count += 1

        conn.commit()
        return {"created": created_count, "updated": updated_count, "warnings": warnings}
    finally:
        conn.close()


def import_cennik_from_excel(filepath, db_path):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header:
        return {"imported": 0, "updated": 0, "errors": ["Brak nagłówków w pliku."]}

    headers = [str(h).strip().lower() if h is not None else "" for h in header]
    idx = {h: i for i, h in enumerate(headers)}
    if "klucz" not in idx:
        for i, h in enumerate(headers):
            if h and "klucz" in h:
                idx["klucz"] = i
                break
    if "klucz" not in idx:
        return {"imported": 0, "updated": 0, "errors": ["Brakuje kolumny 'Klucz'."]}

    conn = sqlite3.connect(db_path)
    cur = conn.cursor()
    imported = 0
    updated = 0
    errors = []

    def _get_cell(row, name_variants):
        for name in name_variants:
            i = idx.get(name)
            if i is not None and i < len(row):
                return row[i]
        return None

    map_variants = {
        "cena_netto_2024": ["cena netto 2024", "cena_netto_2024", "cena netto"],
        "wzrost": ["wzrost ilości dokumentów", "wzrost_ilosci_dokumentow", "wzrost"],
        "cena_za_dokument": ["cena za dokument", "cena_za_dokument"],
        "za_przekroczenie": ["za przekroczenie limitu netto", "za_przekroczenie"],
        "zdrowotna": ["zdrowotna"],
        "netto": ["netto"],
        "brutto": ["brutto"]
    }

    for rownum, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(cell not in (None, "") for cell in row):
            continue
        try:
            key_cell = row[idx["klucz"]]
            if key_cell is None:
                errors.append(f"Wiersz {rownum}: brak klucza")
                continue
            key = str(key_cell).strip()

            def parse_num(val):
                if val is None or (isinstance(val, str) and val.strip() == ""):
                    return None
                try:
                    return float(str(val).replace(",", "."))
                except Exception:
                    return None

            cena_netto_2024 = parse_num(_get_cell(row, map_variants["cena_netto_2024"]))
            wzrost = parse_num(_get_cell(row, map_variants["wzrost"]))
            cena_za_dokument = parse_num(_get_cell(row, map_variants["cena_za_dokument"]))
            za_przekroczenie = parse_num(_get_cell(row, map_variants["za_przekroczenie"]))
            zdrowotna = parse_num(_get_cell(row, map_variants["zdrowotna"]))
            netto = parse_num(_get_cell(row, map_variants["netto"]))
            brutto = parse_num(_get_cell(row, map_variants["brutto"]))

            now = time.strftime("%Y-%m-%d %H:%M:%S")
            cur.execute("SELECT id FROM cenniki WHERE klucz = ?", (key,))
            existing = cur.fetchone()
            if existing:
                cur.execute("""
                    UPDATE cenniki SET
                        cena_netto_2024 = ?, wzrost_ilosci_dokumentow = ?, cena_za_dokument = ?,
                        za_przekroczenie_limitu_netto = ?, zdrowotna = ?, netto = ?, brutto = ?, updated_at = ?
                    WHERE klucz = ?
                """, (cena_netto_2024, wzrost, cena_za_dokument, za_przekroczenie, zdrowotna, netto, brutto, now, key))
                updated += 1
            else:
                cur.execute("""
                    INSERT INTO cenniki (klucz, cena_netto_2024, wzrost_ilosci_dokumentow, cena_za_dokument,
                                         za_przekroczenie_limitu_netto, zdrowotna, netto, brutto, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (key, cena_netto_2024, wzrost, cena_za_dokument, za_przekroczenie, zdrowotna, netto, brutto, now))
                imported += 1
        except Exception as e:
            errors.append(f"Wiersz {rownum}: {str(e)}")
            continue

    conn.commit()
    conn.close()
    return {"imported": imported, "updated": updated, "errors": errors}

